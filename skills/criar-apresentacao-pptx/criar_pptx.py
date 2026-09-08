# -*- coding: utf-8 -*-
"""Gera um .pptx a partir de um roteiro JSON (ver SKILL.md).

  python criar_pptx.py roteiro.json [saida.pptx]

Instala python-pptx automaticamente se faltar.
"""
import json, os, subprocess, sys

def _garante_pptx():
    try:
        import pptx  # noqa: F401
    except ImportError:
        print('instalando python-pptx…', flush=True)
        subprocess.check_call([sys.executable, '-m', 'pip', 'install', '--user', '--quiet', 'python-pptx'])
_garante_pptx()

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

TEMAS = {
    'azul':    {'primaria': '1F4E79', 'secundaria': '2E75B6', 'acento': 'F2B134', 'fundo': 'FFFFFF', 'texto': '1B1B1B', 'suave': 'EAF1FA'},
    'roxo':    {'primaria': '4B2E83', 'secundaria': '6D4DE6', 'acento': 'F4B942', 'fundo': 'FFFFFF', 'texto': '1B1B1B', 'suave': 'F1EDFB'},
    'verde':   {'primaria': '145A32', 'secundaria': '1E8449', 'acento': 'F39C12', 'fundo': 'FFFFFF', 'texto': '1B1B1B', 'suave': 'EAF6EE'},
    'grafite': {'primaria': '2B2B2B', 'secundaria': '555555', 'acento': 'E67E22', 'fundo': 'FFFFFF', 'texto': '1B1B1B', 'suave': 'F0F0F0'},
    'laranja': {'primaria': 'B9470B', 'secundaria': 'E67E22', 'acento': '1F4E79', 'fundo': 'FFFFFF', 'texto': '1B1B1B', 'suave': 'FDF0E6'},
}

def rgb(h):
    return RGBColor.from_string(h)

class Deck:
    def __init__(self, tema, rodape=''):
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self.t = TEMAS.get(tema, TEMAS['azul'])
        self.rodape = rodape
        self.n = 0
        self.W = self.prs.slide_width
        self.H = self.prs.slide_height

    # ---------- primitivas
    def novo(self):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])  # em branco
        self.n += 1
        bg = s.background.fill
        bg.solid(); bg.fore_color.rgb = rgb(self.t['fundo'])
        return s

    def rect(self, s, x, y, w, h, cor, forma=MSO_SHAPE.RECTANGLE):
        r = s.shapes.add_shape(forma, x, y, w, h)
        r.fill.solid(); r.fill.fore_color.rgb = rgb(cor)
        r.line.fill.background()
        r.shadow.inherit = False
        return r

    def texto(self, s, x, y, w, h, txt, tam=18, cor=None, negrito=False, alinh=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, fonte='Calibri'):
        tb = s.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.05)
        p = tf.paragraphs[0]
        p.alignment = alinh
        r = p.add_run(); r.text = str(txt)
        r.font.size = Pt(tam); r.font.bold = negrito; r.font.name = fonte
        r.font.color.rgb = rgb(cor or self.t['texto'])
        return tb

    def topicos(self, s, x, y, w, h, itens, tam=20, cor=None):
        tb = s.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame; tf.word_wrap = True
        first = True
        def add(item, nivel):
            nonlocal first
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.level = nivel
            r = p.add_run(); r.text = ('•  ' if nivel == 0 else '–  ') + str(item)
            r.font.size = Pt(tam - 3 * nivel); r.font.name = 'Calibri'
            r.font.color.rgb = rgb(cor or self.t['texto'])
            p.space_after = Pt(8)
        for it in itens:
            if isinstance(it, list):
                for sub in it: add(sub, 1)
            else:
                add(it, 0)
        return tb

    def cabecalho(self, s, titulo):
        self.rect(s, 0, 0, self.W, Inches(1.1), self.t['primaria'])
        self.rect(s, 0, Inches(1.1), self.W, Inches(0.06), self.t['acento'])
        self.texto(s, Inches(0.6), Inches(0.18), self.W - Inches(1.2), Inches(0.8), titulo, tam=30, cor='FFFFFF', negrito=True, anchor=MSO_ANCHOR.MIDDLE)

    def rodape_num(self, s):
        if self.rodape:
            self.texto(s, Inches(0.6), self.H - Inches(0.5), Inches(8), Inches(0.35), self.rodape, tam=10, cor='888888')
        self.texto(s, self.W - Inches(1.6), self.H - Inches(0.5), Inches(1), Inches(0.35), str(self.n), tam=10, cor='888888', alinh=PP_ALIGN.RIGHT)

    def notas(self, s, txt):
        if txt:
            s.notes_slide.notes_text_frame.text = str(txt)

    # ---------- tipos de slide
    def capa(self, titulo, subtitulo=''):
        s = self.novo()
        self.rect(s, 0, 0, self.W, self.H, self.t['primaria'])
        self.rect(s, 0, Inches(4.9), self.W, Inches(0.08), self.t['acento'])
        self.texto(s, Inches(0.9), Inches(2.2), self.W - Inches(1.8), Inches(2.4), titulo, tam=44, cor='FFFFFF', negrito=True, anchor=MSO_ANCHOR.BOTTOM)
        if subtitulo:
            self.texto(s, Inches(0.9), Inches(5.1), self.W - Inches(1.8), Inches(1.2), subtitulo, tam=20, cor='DDE6F2')
        return s

    def secao(self, sl):
        s = self.novo()
        self.rect(s, 0, 0, Inches(0.5), self.H, self.t['acento'])
        self.texto(s, Inches(1.2), Inches(2.6), self.W - Inches(2.4), Inches(1.6), sl.get('titulo', ''), tam=40, cor=self.t['primaria'], negrito=True, anchor=MSO_ANCHOR.MIDDLE)
        if sl.get('subtitulo'):
            self.texto(s, Inches(1.2), Inches(4.2), self.W - Inches(2.4), Inches(1), sl['subtitulo'], tam=20, cor='666666')
        self.rodape_num(s); self.notas(s, sl.get('notas'))

    def slide_topicos(self, sl):
        s = self.novo()
        self.cabecalho(s, sl.get('titulo', ''))
        self.topicos(s, Inches(0.8), Inches(1.5), self.W - Inches(1.6), self.H - Inches(2.3), sl.get('itens', []))
        self.rodape_num(s); self.notas(s, sl.get('notas'))

    def duas_colunas(self, sl):
        s = self.novo()
        self.cabecalho(s, sl.get('titulo', ''))
        colw = (self.W - Inches(1.9)) / 2
        for i, lado in enumerate(('esquerda', 'direita')):
            col = sl.get(lado, {})
            x = Inches(0.8) + i * (colw + Inches(0.3))
            self.rect(s, x, Inches(1.5), colw, Inches(0.55), self.t['suave'])
            self.texto(s, x + Inches(0.1), Inches(1.5), colw, Inches(0.55), col.get('titulo', ''), tam=18, cor=self.t['primaria'], negrito=True, anchor=MSO_ANCHOR.MIDDLE)
            self.topicos(s, x, Inches(2.2), colw, self.H - Inches(3), col.get('itens', []), tam=18)
        self.rodape_num(s); self.notas(s, sl.get('notas'))

    def tabela(self, sl):
        s = self.novo()
        self.cabecalho(s, sl.get('titulo', ''))
        cols = sl.get('colunas', []); rows = sl.get('linhas', [])
        if not cols:
            raise ValueError('slide tabela sem "colunas"')
        shape = s.shapes.add_table(len(rows) + 1, len(cols), Inches(0.8), Inches(1.6), self.W - Inches(1.6), Inches(0.4) * (len(rows) + 1))
        tbl = shape.table
        for c, nome in enumerate(cols):
            cell = tbl.cell(0, c); cell.text = str(nome)
            cell.fill.solid(); cell.fill.fore_color.rgb = rgb(self.t['primaria'])
            for p in cell.text_frame.paragraphs:
                for r in p.runs: r.font.bold = True; r.font.size = Pt(14); r.font.color.rgb = rgb('FFFFFF')
        for ri, row in enumerate(rows, start=1):
            for c in range(len(cols)):
                cell = tbl.cell(ri, c); cell.text = str(row[c]) if c < len(row) else ''
                cell.fill.solid(); cell.fill.fore_color.rgb = rgb(self.t['suave'] if ri % 2 else 'FFFFFF')
                for p in cell.text_frame.paragraphs:
                    for r in p.runs: r.font.size = Pt(13); r.font.color.rgb = rgb(self.t['texto'])
        self.rodape_num(s); self.notas(s, sl.get('notas'))

    def imagem(self, sl):
        s = self.novo()
        self.cabecalho(s, sl.get('titulo', ''))
        caminho = sl.get('caminho', '')
        if not os.path.exists(caminho):
            raise FileNotFoundError(f'imagem não encontrada: {caminho}')
        maxw, maxh = self.W - Inches(1.6), self.H - Inches(2.6)
        pic = s.shapes.add_picture(caminho, Inches(0.8), Inches(1.5))
        ratio = min(maxw / pic.width, maxh / pic.height, 1.0) if pic.width and pic.height else 1.0
        pic.width = int(pic.width * ratio); pic.height = int(pic.height * ratio)
        pic.left = int((self.W - pic.width) / 2); pic.top = Inches(1.5)
        if sl.get('legenda'):
            self.texto(s, Inches(0.8), self.H - Inches(1.05), self.W - Inches(1.6), Inches(0.4), sl['legenda'], tam=12, cor='666666', alinh=PP_ALIGN.CENTER)
        self.rodape_num(s); self.notas(s, sl.get('notas'))

    def citacao(self, sl):
        s = self.novo()
        self.rect(s, 0, 0, self.W, self.H, self.t['suave'])
        self.texto(s, Inches(1.2), Inches(1.2), Inches(1.5), Inches(1.5), '“', tam=120, cor=self.t['acento'], negrito=True)
        self.texto(s, Inches(1.5), Inches(2.4), self.W - Inches(3), Inches(2.6), sl.get('texto', ''), tam=30, cor=self.t['primaria'], anchor=MSO_ANCHOR.MIDDLE)
        if sl.get('autor'):
            self.texto(s, Inches(1.5), Inches(5.2), self.W - Inches(3), Inches(0.6), '— ' + sl['autor'], tam=18, cor='555555')
        self.rodape_num(s); self.notas(s, sl.get('notas'))

    def numeros(self, sl):
        s = self.novo()
        self.cabecalho(s, sl.get('titulo', ''))
        itens = sl.get('itens', [])[:4]
        if not itens:
            raise ValueError('slide numeros sem "itens"')
        gap = Inches(0.3)
        w = (self.W - Inches(1.6) - gap * (len(itens) - 1)) / len(itens)
        for i, it in enumerate(itens):
            x = Inches(0.8) + i * (w + gap)
            self.rect(s, x, Inches(2.2), w, Inches(2.8), self.t['suave'], MSO_SHAPE.ROUNDED_RECTANGLE)
            self.texto(s, x, Inches(2.5), w, Inches(1.3), it.get('valor', ''), tam=40, cor=self.t['secundaria'], negrito=True, alinh=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            self.texto(s, x + Inches(0.2), Inches(3.9), w - Inches(0.4), Inches(0.9), it.get('rotulo', ''), tam=16, cor='444444', alinh=PP_ALIGN.CENTER)
        self.rodape_num(s); self.notas(s, sl.get('notas'))

    def encerramento(self, sl):
        s = self.novo()
        self.rect(s, 0, 0, self.W, self.H, self.t['primaria'])
        self.texto(s, Inches(0.9), Inches(0.9), self.W - Inches(1.8), Inches(1.2), sl.get('titulo', 'Obrigado'), tam=40, cor='FFFFFF', negrito=True)
        if sl.get('itens'):
            self.topicos(s, Inches(0.9), Inches(2.3), self.W - Inches(1.8), Inches(3.2), sl['itens'], tam=22, cor='FFFFFF')
        if sl.get('contato'):
            self.texto(s, Inches(0.9), self.H - Inches(1.2), self.W - Inches(1.8), Inches(0.6), sl['contato'], tam=16, cor='DDE6F2')
        self.notas(s, sl.get('notas'))

def main():
    if len(sys.argv) < 2:
        sys.exit('uso: python criar_pptx.py roteiro.json [saida.pptx]')
    with open(sys.argv[1], encoding='utf-8-sig') as f:
        rot = json.load(f)
    saida = sys.argv[2] if len(sys.argv) > 2 else os.path.splitext(sys.argv[1])[0] + '.pptx'
    d = Deck(rot.get('tema', 'azul'), rot.get('rodape', ''))
    d.capa(rot.get('titulo', 'Apresentação'), rot.get('subtitulo', ''))
    tipos = {'secao': d.secao, 'topicos': d.slide_topicos, 'duas_colunas': d.duas_colunas, 'tabela': d.tabela,
             'imagem': d.imagem, 'citacao': d.citacao, 'numeros': d.numeros, 'encerramento': d.encerramento}
    for i, sl in enumerate(rot.get('slides', []), start=1):
        tipo = sl.get('tipo', 'topicos')
        if tipo == 'capa':
            continue
        if tipo not in tipos:
            sys.exit(f'slide {i}: tipo desconhecido "{tipo}". Use: {", ".join(tipos)}')
        try:
            tipos[tipo](sl)
        except Exception as e:
            sys.exit(f'slide {i} ({tipo}): {e}')
    os.makedirs(os.path.dirname(os.path.abspath(saida)), exist_ok=True)
    d.prs.save(saida)
    print(f'OK: {os.path.abspath(saida)} ({d.n} slides)')

if __name__ == '__main__':
    main()
